"""docs/analiz/Seviye_Planlama.xlsx (GECE_2026-10-13 B6, PM 2026-10-13 ~23:30): the difficulty
model as a spreadsheet Zeyd can play with.

  Parametre  the genre table (s, P, bomb margin, nail shares, first-try and near-miss targets)
  Seviye     the 40 web levels, prefilled; formulas give s, the timer, fuse, nail counters, the
             median player's time left and stars
  Fabrika    per slot of the redesigned set: the approved level or the best finalist (redesign.json)

Formulas follow Zeyd's style: LET, one line, no spaces, English names (Excel shows them with ';'),
XLOOKUP and CEILING.MATH. In the file they are stored the way Excel stores them (_xlfn. / _xlpm.).
The workbook is marked to recalculate when it opens. The same rules are also computed here and
compared with DifficultyModel's own suggestions (the 40-level report probe); any difference is
printed (see main()).

usage: python Tools/make_planning_xlsx.py <web40.tsv> <levels.jsonl from the 40-level report probe>
needs openpyxl (pip install openpyxl)."""
import json, math, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
OUT = os.path.join(ROOT, 'docs', 'analiz', 'Seviye_Planlama.xlsx')

# PRODUCT.md "Zorluk modeli (başlangıç değerleri)" + "Keyif eşikleri (kontrol noktası 3/4)"
GENRES = [  # name, s, P, bomb margin, nail share 1, nail share 2, first-try target, near-miss target
    ('Tanıtım', 1.1, 2.2, 10, .225, .225, '≥ %95', '%0–5'),
    ('Nefes', 1.0, 2.0, 8, .225, .225, '≥ %90', '%0–5'),
    ('Kolay', 1.0, 2.0, 8, .225, .225, '%85–95', '%0–5'),
    ('Orta', 1.1, 1.6, 8, .30, .30, '%65–85', '%5–15'),
    ('Zor', 1.2, 1.3, 5, .275, .375, '%40–60', '%15–30'),
    ('Çok zor', 1.3, 1.15, 2, .275, .375, '%20–35', '%30–50'),
]
G = {g[0]: g for g in GENRES}


def xl(f):
    """Excel's stored form: _xlfn. for newer functions, _xlpm. for LET names."""
    import re
    f = f.replace('XLOOKUP(', '_xlfn.XLOOKUP(').replace('CEILING.MATH(', '_xlfn.CEILING.MATH(').replace('LET(', '_xlfn.LET(')
    m = re.search(r'_xlfn\.LET\((.*)\)$', f)
    if m:
        names = set(re.findall(r'(?:\(|,)([a-z])(?=,)', '(' + m.group(1)))
        for n in sorted(names, key=len, reverse=True):
            f = re.sub(r'(?<![A-Za-z_.$!])' + n + r'(?![A-Za-z0-9_(])', '_xlpm.' + n, f)
    return '=' + f


def ceil5(x):
    return int(math.ceil(x / 5 - 1e-9) * 5)


def main():
    web = [l.rstrip('\n').split('\t') for l in open(sys.argv[1], encoding='utf-8') if l.strip()]
    rep = {json.loads(l)['id']: json.loads(l) for l in open(sys.argv[2], encoding='utf-8') if l.startswith('{')}
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1F4E57')

    # ---- Parametre ----
    ps = wb.active
    ps.title = 'Parametre'
    cols = ['Tür', 's (sn/parça)', 'P (süre payı)', 'Bomba payı', 'Çivi payı 1', 'Çivi payı 2', 'İlk deneme hedefi', 'Near-miss hedefi']
    ps.append(cols)
    for g in GENRES: ps.append(list(g))
    ps.append([])
    ps.append(['Kaynak: PRODUCT.md "Zorluk modeli (başlangıç değerleri)". Süre = ⌈gövde × s × P⌉₅, en fazla 75 sn; s yapıdan: 0,6 + 0,16 × hamle başına okuma (+0,1 bombalı, ilk bölümde 1,35).'])
    ps.append(['Fitil = en kısa kazma yolu + bomba payı. Çivi sayacı = gövde × pay (üst sınır %40). Medyan oyuncunun kalan süresi = 1 − gövde × s / süre; yıldız: ≥ %45 → 3, ≥ %20 → 2, altı 1.'])

    # ---- Seviye ----
    ls = wb.create_sheet('Seviye')
    hdr = ['Seviye', 'Tür', 'Gövde', 'Hamle başına okuma', 'Bomba', 'En kısa kazma (en yakın bomba)', 'İlk bölüm', 'Çivi', 'Mevcut süre',
           's', 'P', 'Önerilen süre', 'Önerilen fitil', 'Çivi sayacı 1', 'Çivi sayacı 2', 'Medyan kalan süre', 'Medyan yıldız']
    ls.append(hdr)
    mism = []
    for i, w in enumerate(web, start=2):
        lid, g, read, bombs, dig, exact, nails, timer = w[0], int(w[1]), float(w[2]), int(w[3]), int(w[4]), w[5], int(w[6]), int(w[7])
        genre = rep[lid]['genre']
        first = 1 if lid == 'L01' else 0
        tg = G[genre]
        s = 1.35 if first else .6 + .16 * read + (.1 if bombs else 0)
        t = min(75, ceil5(g * s * tg[2]))
        fuse = dig + tg[3] if bombs else ''
        n1 = max(1, int(math.floor(g * tg[4] + .5))) if nails >= 1 else ''
        n2 = max(1, int(math.floor(g * tg[5] + .5))) if nails >= 2 else ''
        left = 1 - g * s / t
        stars = 3 if left >= .45 else 2 if left >= .2 else 1
        # cross-check with DifficultyModel (the 40-level report probe)
        if rep[lid]['sugg'] != t: mism.append(f'{lid} süre {t} ≠ {rep[lid]["sugg"]}')
        if abs(rep[lid]['s'] - round(s, 3)) > .0015: mism.append(f'{lid} s {s:.3f} ≠ {rep[lid]["s"]}')
        if bombs and rep[lid]['fuses']:
            want = min(int(x.split('>')[1].rstrip('~')) for x in rep[lid]['fuses'].split('/'))   # the nearest bomb
            if fuse != want and exact == '1': mism.append(f'{lid} fitil {fuse} ≠ {want}')
        r = str(i)
        L = 'Parametre!$A$2:$A$7'
        row = [lid, genre, g, round(read, 3), bombs, dig if bombs else '', first, nails, timer,
               xl(f'LET(r,D{r},b,E{r},f,G{r},IF(f=1,1.35,0.6+0.16*r+IF(b>0,0.1,0)))'),
               xl(f'XLOOKUP(B{r},{L},Parametre!$C$2:$C$7)'),
               xl(f'LET(g,C{r},s,J{r},p,K{r},MIN(75,CEILING.MATH(g*s*p,5)))'),
               xl(f'IF(E{r}=0,"",F{r}+XLOOKUP(B{r},{L},Parametre!$D$2:$D$7))'),
               xl(f'LET(g,C{r},k,XLOOKUP(B{r},{L},Parametre!$E$2:$E$7),IF(H{r}<1,"",MAX(1,ROUND(g*k,0))))'),
               xl(f'LET(g,C{r},k,XLOOKUP(B{r},{L},Parametre!$F$2:$F$7),IF(H{r}<2,"",MAX(1,ROUND(g*k,0))))'),
               xl(f'LET(g,C{r},s,J{r},t,L{r},1-g*s/t)'),
               xl(f'LET(m,P{r},IF(m>=0.45,3,IF(m>=0.2,2,1)))')]
        ls.append(row)
    # ---- Fabrika ----
    fs = wb.create_sheet('Fabrika')
    fs.append(['Slot', 'Tür', 'Yeni', 'Seviye', 'Durum', 'Tahta', 'Gövde', 'Süre', 'P', 'Okuma', 'İlk hamlede serbest',
               'Keyif tahmini', 'Hayal kırıklığı', 'Sıkıcı', 'İlk deneme', '3 denemede', 'Near-miss', 'Eğri uyumu'])
    red = json.load(open(os.path.join(ROOT, 'Assets/_Game/Levels/redesign.json'), encoding='utf-8'))
    plan = json.load(open(os.path.join(ROOT, 'Assets/_Game/Levels/fabrika_plan.json'), encoding='utf-8'))['slots']
    label = {'tanitim': 'Tanıtım', 'nefes': 'Nefes', 'kolay': 'Kolay', 'orta': 'Orta', 'zor': 'Zor', 'cokzor': 'Çok zor'}
    new = {'nail': 'çivi', 'scissors': 'makas', 'seal': 'mühür', 'wand': 'değnek', 'chain': 'zincir', 'clock': 'saat', 'bomb': 'bomba', 'hammer': 'çekiç'}
    for sp in plan:
        es = [l for l in red['levels'] if l['fabrika']['slot'] == sp['slot']]
        pick = next((l for l in es if l['fabrika']['status'] == 'onayli'), None) or \
            max((l for l in es if l['fabrika']['status'] == 'finalist'), key=lambda l: l['fabrika']['scores'].get('keyif', 0) - 2 * l['fabrika']['scores'].get('hayalKirikligi', 0), default=None)
        if pick is None:
            fs.append([sp['slot'], label[sp['genre']], new.get(sp['new'], ''), '—', 'aday yok']); continue
        sc = pick['fabrika']['scores']
        fs.append([sp['slot'], label[sp['genre']], new.get(sp['new'], ''), pick['id'], pick['fabrika']['status'], f"{pick['w']}×{pick['h']}",
                   sc.get('govde'), pick['timer'], sc.get('p'), sc.get('okuma'), sc.get('k0'), sc.get('keyif'), sc.get('hayalKirikligi'),
                   sc.get('sikici'), sc.get('ilkDeneme'), sc.get('ucDeneme'), sc.get('nearMiss'), sc.get('uyum')])
    for row in fs.iter_rows(min_row=2):
        for c in row[11:17]:
            if isinstance(c.value, (int, float)): c.number_format = '0%'

    for sh in (ps, ls, fs):
        for c in sh[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(wrap_text=True, vertical='top')
        sh.freeze_panes = 'A2'
        for i in range(1, sh.max_column + 1): sh.column_dimensions[get_column_letter(i)].width = 14
    for row in ls.iter_rows(min_row=2, min_col=16, max_col=16):
        for c in row: c.number_format = '0%'
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(OUT)
    print(f'{OUT}: Seviye {len(web)} satır, Fabrika {len(plan)} slot; DifficultyModel ile uyuşmazlık: {len(mism)}')
    for m in mism: print('  ' + m)


if __name__ == '__main__':
    main()
